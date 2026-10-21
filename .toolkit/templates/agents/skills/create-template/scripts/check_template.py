"""Prueft ein fertiges Template (.md oder .docx) auf verbliebene Quelldaten und listet Platzhalter.

Aufruf: python check_template.py <template.md|template.docx> [--forbid "Begriff" ...] [--forbid-file liste.txt]

- --forbid / --forbid-file: Begriffe der Quelle, die NICHT im Template bleiben duerfen
  (case-insensitiv; Datei = ein Begriff pro Zeile). Treffer -> Exit-Code 1.
- Warnungen (kein Fehler): verdaechtige Konkretwerte ausserhalb von Platzhaltern
  (E-Mail, Telefonnummer, Datum, Betrag/Waehrung, Zahlen ab 4 Stellen).
- Meldet Anzahl und Liste der Platzhalter ([...], {{...}}, XXXX).
"""
import re
import sys

PLACEHOLDER = re.compile(r"\[[^\]]+\]|\{\{[^}]+\}\}|X{3,}")
SUSPICIOUS = {
    "E-Mail": re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+"),
    "Telefon": re.compile(r"(?<!\d)(?:\+\d{2}|0\d{2})[\s./]?\d{2,3}[\s./]?\d{2,3}[\s./]?\d{2}(?!\d)"),
    "Datum": re.compile(r"\b\d{1,2}\.\d{1,2}\.\d{2,4}\b|\b\d{4}-\d{2}-\d{2}\b"),
    "Betrag": re.compile(r"\b(?:CHF|EUR|USD|Fr\.)\s?\d[\d'.,]*|\d[\d'.,]*\s?(?:CHF|EUR|USD)\b"),
    "Zahl >= 4 Stellen": re.compile(r"(?<![\w.])\d{4,}(?![\w.])"),
}


def read_lines(path):
    if path.lower().endswith(".docx"):
        from docx import Document

        doc = Document(path)
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.extend(c.text for c in row.cells)
        for section in doc.sections:
            for part in (section.header, section.footer):
                lines.extend(p.text for p in part.paragraphs)
        return [l for l in lines if l.strip()]
    with open(path, encoding="utf-8-sig") as f:
        return [l.rstrip("\n") for l in f if l.strip()]


def parse_args(argv):
    path, forbid, i = None, [], 0
    while i < len(argv):
        if argv[i] == "--forbid" and i + 1 < len(argv):
            forbid.append(argv[i + 1])
            i += 2
        elif argv[i] == "--forbid-file" and i + 1 < len(argv):
            with open(argv[i + 1], encoding="utf-8-sig") as f:
                forbid.extend(l.strip() for l in f if l.strip())
            i += 2
        elif path is None:
            path = argv[i]
            i += 1
        else:
            sys.exit(__doc__)
    if path is None:
        sys.exit(__doc__)
    return path, forbid


def main(argv):
    sys.stdout.reconfigure(encoding="utf-8")
    path, forbid = parse_args(argv)
    lines = read_lines(path)

    errors = 0
    print("=== VERBOTENE BEGRIFFE ===")
    for term in forbid:
        hits = [l for l in lines if term.lower() in l.lower()]
        for h in hits:
            print(f"TREFFER '{term}': {h.strip()[:140]}")
        errors += len(hits)
    if not forbid:
        print("(keine Begriffsliste angegeben)")
    elif errors == 0:
        print("OK - keine Treffer")

    print("\n=== WARNUNGEN (Konkretwerte ausserhalb von Platzhaltern) ===")
    warnings = 0
    for line in lines:
        stripped = PLACEHOLDER.sub("", line)
        for label, pattern in SUSPICIOUS.items():
            for m in pattern.finditer(stripped):
                print(f"{label}: '{m.group(0)}' in: {line.strip()[:120]}")
                warnings += 1
    if warnings == 0:
        print("keine")

    print("\n=== PLATZHALTER ===")
    found = [m.group(0) for l in lines for m in PLACEHOLDER.finditer(l)]
    print(f"Anzahl: {len(found)} (eindeutig: {len(set(found))})")
    for p in sorted(set(found)):
        print(f"  {p}")
    if not found:
        print("WARNUNG: keine Platzhalter gefunden - ist das Dokument wirklich ein Template?")

    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main(sys.argv[1:])
