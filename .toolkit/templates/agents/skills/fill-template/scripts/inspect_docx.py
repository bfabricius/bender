"""Gibt Struktur, Formatvorlagen und Platzhalter eines .docx aus (Template-Analyse).

Aufruf: python inspect_docx.py <datei.docx>
"""
import re
import sys

from docx import Document

PLACEHOLDER = re.compile(r"\[[^\]]+\]|\{\{[^}]+\}\}|X{3,}")


def main(path):
    sys.stdout.reconfigure(encoding="utf-8")
    doc = Document(path)

    print("=== PARAGRAPHEN (Index | Style | Text) ===")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"{i} | {p.style.name} | {p.text}")

    print("\n=== TABELLEN ===")
    for ti, table in enumerate(doc.tables):
        style = table.style.name if table.style is not None else None
        print(f"TABLE {ti}: {len(table.rows)}x{len(table.columns)} style={style!r}")
        for ri, row in enumerate(table.rows):
            print(f"  r{ri}: " + " | ".join(c.text.replace("\n", " / ") for c in row.cells))

    print("\n=== KOPF-/FUSSZEILE ===")
    for si, section in enumerate(doc.sections):
        for label, part in (("header", section.header), ("footer", section.footer)):
            for p in part.paragraphs:
                if p.text.strip():
                    print(f"section {si} {label}: {p.text}")

    print("\n=== PLATZHALTER ===")
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            texts.extend(c.text for c in row.cells)
    for t in texts:
        for m in PLACEHOLDER.findall(t):
            print(m)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
