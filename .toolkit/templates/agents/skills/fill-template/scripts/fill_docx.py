"""Befüllt eine Kopie eines .docx-Templates anhand einer JSON-Operationsliste.

Aufruf: python fill_docx.py <template.docx> <ausgabe.docx> <ops.json>

ops.json = Liste von Operationen. Paragraph-Indizes beziehen sich auf doc.paragraphs des
ORIGINAL-Templates (siehe inspect_docx.py) und bleiben trotz Einfügen/Löschen gültig.

  {"op": "set_paragraph", "index": 12, "text": "Neuer Text"}
  {"op": "insert_after",  "index": 12, "text": "Zusatzzeile"}      # klont Format von Paragraph 12
  {"op": "insert_after",  "index": 12, "text": "...", "style": "List Bullet"}
  {"op": "delete_paragraph", "index": 13}
  {"op": "set_cell", "table": 0, "row": 1, "col": 1, "text": "Wert"}
  {"op": "replace_text", "find": "[Projektname]", "replace": "Mandat Mustermandant"}

set_* behält das Format des ersten Runs. replace_text wirkt auf Body, Tabellen, Kopf-/Fusszeilen
(Run-übergreifende Treffer werden im ersten Run des Paragraphen zusammengeführt).
Mehrere insert_after auf denselben Index erscheinen in der angegebenen Reihenfolge.
"""
import copy
import json
import sys

from docx import Document
from docx.text.paragraph import Paragraph


def set_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def replace_in_paragraph(paragraph, find, replace):
    if find not in paragraph.text:
        return
    for r in paragraph.runs:
        if find in r.text:
            r.text = r.text.replace(find, replace)
            return
    set_text(paragraph, paragraph.text.replace(find, replace))


def iter_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def main(template, output, ops_path):
    doc = Document(template)
    with open(ops_path, encoding="utf-8-sig") as f:
        ops = json.load(f)

    original = list(doc.paragraphs)
    last_inserted = {}

    for op in ops:
        kind = op["op"]
        if kind == "set_paragraph":
            set_text(original[op["index"]], op["text"])
        elif kind == "insert_after":
            anchor = original[op["index"]]
            new_el = copy.deepcopy(anchor._element)
            last_el = last_inserted.get(op["index"], anchor._element)
            last_el.addnext(new_el)
            last_inserted[op["index"]] = new_el
            new_p = Paragraph(new_el, anchor._parent)
            set_text(new_p, op["text"])
            if "style" in op:
                new_p.style = doc.styles[op["style"]]
        elif kind == "delete_paragraph":
            el = original[op["index"]]._element
            el.getparent().remove(el)
        elif kind == "set_cell":
            cell = doc.tables[op["table"]].cell(op["row"], op["col"])
            paragraphs = cell.paragraphs
            set_text(paragraphs[0], op["text"])
            for p in paragraphs[1:]:
                p._element.getparent().remove(p._element)
        elif kind == "replace_text":
            for p in iter_paragraphs(doc):
                replace_in_paragraph(p, op["find"], op["replace"])
        else:
            sys.exit(f"Unbekannte Operation: {kind}")

    doc.save(output)
    print(f"Gespeichert: {output}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(*sys.argv[1:])
